"""
Build output/Primary_School_Classroom_Report_2026.xlsx — exactly the output the
Data Transformation Brief asks for: primary grades P1-P6 only, section 3a
(Per-Grade Breakdown) and section 3b (School-Level Summary).

Unlike the full analysis workbook, the school level here follows the brief's
step 4 literally: it sums the grade rows (SUMIFS on the Per-Grade sheet).

Run:  python scripts/build_brief_workbook.py
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment

import analysis as A
from build_workbook import (
    BOX, F_BOLD, F_INPUT, F_SUB, FILL_INPUT, GAP, NUM, WRAP, gap_colours, header, put, q, status_colours, title,
)

OUT = A.ROOT / "output" / "Primary_School_Classroom_Report_2026.xlsx"
GRADE, SCHOOL, NOTES = "Per-Grade Breakdown", "School-Level Summary", "Notes"
CAPACITY_CELL = f"{q(NOTES)}$B$4"

OUTPUT_COLS = ["Total Students", "Total Classrooms", "Total Classrooms in Double Shift", "Total Classrooms Available",
               "Standard Classroom Capacity", "Required Classrooms", "Gap (Deficit vs Surplus)"]


def build_notes(ws, roster_primary, g_last: int, s_last: int) -> None:
    title(ws, "Notes", "Primary School Infrastructure Analysis 2026 — built from Data_Transformation_Brief.docx.")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    put(ws, 4, 1, "Standard Classroom Capacity", font=F_BOLD)
    c = put(ws, 4, 2, 45, NUM, F_INPUT)
    c.fill = FILL_INPUT
    c.alignment = Alignment(horizontal="left")
    c.comment = Comment("Brief section 4: fixed at 45 students per classroom. Every capacity cell links here.", "Analysis")

    rows = [
        ("Scope", "Primary grades P1-P6 only (brief section 5, step 1)."),
        ("Per-Grade Breakdown", "Brief section 3a: one row per school x grade."),
        ("School-Level Summary", "Brief section 3b: one row per school, P1-P6 summed from the Per-Grade Breakdown."),
        ("Total Classrooms", "Count of classroom rows (class groups) per school and grade."),
        ("Total Classrooms in Double Shift", "The source has no morning/afternoon field. A double-shift classroom is one "
                                             "physical room (same room ID, test_code) used by more than one class group; each "
                                             "extra group counts once. Example: 9 class groups in 6 rooms, 3 of them shared "
                                             "-> 3 double shift, 6 available. Rooms recorded with 3+ groups are split into "
                                             "one room per group in the source (data-entry errors)."),
        ("Total Classrooms Available", "Total Classrooms - Total Classrooms in Double Shift."),
        ("Required Classrooms", "CEILING(Total Students / 45), always rounded up."),
        ("Gap (Deficit vs Surplus)", "Available - Required. Negative = deficit, 0 = exact fit, positive = surplus."),
        ("Rooms shared by two grades", "Counted in each grade that uses them (brief step 4 sums the grades), so a few "
                                       "schools show slightly more available classrooms than physical rooms."),
        ("Source data kept as recorded", "Two rows are graded P1 but labelled as pre-primary groups (N3A, N3); they are "
                                         "included as recorded. See the full analysis workbook for the data-quality list."),
    ]
    r = 6
    ws.cell(row=r, column=1, value="Definitions").font = F_SUB
    for label, text in rows:
        r += 1
        put(ws, r, 1, label, font=F_BOLD).alignment = WRAP
        put(ws, r, 2, text).alignment = WRAP

    r += 2
    ws.cell(row=r, column=1, value="Checks").font = F_SUB
    checks = [
        ("Students (P1-P6) in source", int(roster_primary[A.STUDENTS].sum()),
         f"=SUM({q(GRADE)}D5:D{g_last})", f"=SUM({q(SCHOOL)}C5:C{s_last})"),
        ("Classroom rows (P1-P6) in source", len(roster_primary),
         f"=SUM({q(GRADE)}E5:E{g_last})", f"=SUM({q(SCHOOL)}D5:D{s_last})"),
    ]
    r += 1
    for col, label in enumerate(["Check", "Source", "Per-Grade", "School-Level", "Result"], 1):
        put(ws, r, col, label, font=F_BOLD).border = BOX
    for col in "CDE":
        ws.column_dimensions[col].width = 14
    for label, raw, fg, fs in checks:
        r += 1
        put(ws, r, 1, label, font=F_BOLD)
        put(ws, r, 2, raw, NUM).comment = Comment("Counted from the source roster when the file was built.", "Analysis")
        put(ws, r, 3, fg, NUM)
        put(ws, r, 4, fs, NUM)
        put(ws, r, 5, f'=IF(AND(B{r}=C{r},B{r}=D{r}),"OK","MISMATCH")', font=F_BOLD)
    ws.freeze_panes = "A5"


def build_grade(ws, g) -> int:
    title(ws, "3a. Per-Grade Breakdown", "One row per School x Grade, primary grades P1-P6.")
    header(ws, 4, ["School Code", "School Name", "Grade", *OUTPUT_COLS, "Status"],
           [12, 40, 9, 14, 17, 20.5, 18, 17, 17, 18, 12])
    for i, row in enumerate(g.itertuples(index=False)):
        r = 5 + i
        put(ws, r, 1, int(row.school_code))
        put(ws, r, 2, row.school_name)
        put(ws, r, 3, row.grade)
        put(ws, r, 4, int(row.total_students), NUM)
        put(ws, r, 5, int(row.total_classrooms), NUM)
        put(ws, r, 6, int(row.double_shift), NUM)
        put(ws, r, 7, f"=E{r}-F{r}", NUM)
        put(ws, r, 8, f"={CAPACITY_CELL}", NUM)
        put(ws, r, 9, f"=ROUNDUP(D{r}/H{r},0)", NUM)
        put(ws, r, 10, f"=G{r}-I{r}", GAP)
        put(ws, r, 11, f'=IF(J{r}<0,"Deficit",IF(J{r}=0,"Exact fit","Surplus"))')
    last = 4 + len(g)
    gap_colours(ws, f"J5:J{last}")
    status_colours(ws, f"K5:K{last}", "$K5")
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:K{last}"
    return last


def build_school(ws, schools, g_last: int) -> int:
    title(ws, "3b. School-Level Summary", "One row per school, grades P1-P6 combined (sums of the Per-Grade Breakdown).")
    header(ws, 4, ["School Code", "School Name", *OUTPUT_COLS, "Status"], [12, 40, 14, 17, 20.5, 18, 17, 17, 18, 12])
    key = f"{q(GRADE)}$A$5:$A${g_last}"
    col = lambda c: f"{q(GRADE)}${c}$5:${c}${g_last}"  # noqa: E731
    for i, row in enumerate(schools.itertuples(index=False)):
        r = 5 + i
        put(ws, r, 1, int(row.school_code))
        put(ws, r, 2, row.school_name)
        put(ws, r, 3, f"=SUMIFS({col('D')},{key},$A{r})", NUM)
        put(ws, r, 4, f"=SUMIFS({col('E')},{key},$A{r})", NUM)
        put(ws, r, 5, f"=SUMIFS({col('F')},{key},$A{r})", NUM)
        put(ws, r, 6, f"=D{r}-E{r}", NUM)
        put(ws, r, 7, f"={CAPACITY_CELL}", NUM)
        put(ws, r, 8, f"=ROUNDUP(C{r}/G{r},0)", NUM)
        put(ws, r, 9, f"=F{r}-H{r}", GAP)
        put(ws, r, 10, f'=IF(I{r}<0,"Deficit",IF(I{r}=0,"Exact fit","Surplus"))')
    last = 4 + len(schools)
    gap_colours(ws, f"I5:I{last}")
    status_colours(ws, f"J5:J{last}", "$J5")
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:J{last}"
    return last


def main() -> None:
    roster = A.load_roster()
    primary = roster[roster["level"] == "Primary"]
    g = A.grade_table(primary)
    g["_grd"] = g["grade"].map({x: i for i, x in enumerate(A.GRADE_ORDER)})
    g = g.sort_values(["school_code", "_grd"]).reset_index(drop=True)
    schools = g[["school_code", "school_name"]].drop_duplicates("school_code").reset_index(drop=True)

    wb = Workbook()
    ws_grade = wb.active
    ws_grade.title = GRADE
    ws_school = wb.create_sheet(SCHOOL)
    ws_notes = wb.create_sheet(NOTES)

    g_last = build_grade(ws_grade, g)
    s_last = build_school(ws_school, schools, g_last)
    build_notes(ws_notes, primary, g_last, s_last)
    ws_notes.sheet_properties.tabColor = "FFC000"
    wb.calculation.fullCalcOnLoad = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Saved {OUT}  (grade rows {len(g):,}, schools {len(schools):,})")


if __name__ == "__main__":
    main()
